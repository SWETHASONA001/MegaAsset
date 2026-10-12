"""Test fixtures: isolated database + small raw-data tree (real CMRL GTFS + generated documents).

Environment variables are set before the app is imported so settings/engine point at the temp location.
"""
import os
import shutil
import sys
import tempfile
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
PROJECT = BACKEND.parent
TMP = Path(tempfile.mkdtemp(prefix="megaasset_test_"))
RAW = TMP / "raw"

os.environ.update({
    "MEGAASSET_DATABASE_PATH": str(TMP / "test.db"),
    "MEGAASSET_RAW_DATA_DIR": str(RAW),
    "MEGAASSET_UPLOADS_DIR": str(TMP / "uploads"),
    "MEGAASSET_PROCESSED_DIR": str(TMP / "processed"),
    "MEGAASSET_MODELS_DIR": str(BACKEND / "data" / "models"),
    "MEGAASSET_AUTO_INGEST_ON_STARTUP": "false",
    "MEGAASSET_INGEST_WORKERS": "1",
    "MEGAASSET_ENABLE_OCR": "false",
    "MEGAASSET_JWT_SECRET": "test-secret-key-for-megaasset-tests-only",
})
sys.path.insert(0, str(BACKEND))


def _build_raw():
    shutil.copytree(PROJECT / "data" / "raw" / "cmrl" / "gtfs", RAW / "cmrl" / "gtfs")
    docs = RAW / "cmrl" / "documents" / "maintenance"
    docs.mkdir(parents=True)
    import pymupdf
    pdf = pymupdf.open()
    p1 = pdf.new_page()
    p1.insert_text((72, 72), "Escalator maintenance requirements", fontsize=14)
    p1.insert_textbox(pymupdf.Rect(72, 100, 520, 400),
                      "The contractor shall inspect every escalator at Egmore Metro station monthly. "
                      "Escalator breakdowns must be attended within 60 minutes. Preventive maintenance of handrails "
                      "and step chains is mandatory under this contract.", fontsize=11)
    p2 = pdf.new_page()
    p2.insert_textbox(pymupdf.Rect(72, 72, 520, 400),
                      "Lift maintenance: the lifts at Guindy Metro station shall achieve a mean time to repair of 45 minutes. "
                      "Fire safety audits are conducted annually across all stations.", fontsize=11)
    pdf.save(docs / "Test_Escalator_AMC.pdf")
    (docs / "notes.txt").write_text("Housekeeping schedule\n\nPlatforms are cleaned every two hours by the housekeeping contractor.", encoding="utf-8")
    shutil.copy(docs / "notes.txt", docs / "notes_copy.txt")  # duplicate content
    rid = RAW / "cmrl" / "ridership"
    rid.mkdir(parents=True)
    (rid / "cmrl_monthly_ridership.csv").write_text(
        "_id,Month,Closed Loop (Ridership),Closed Loop %,QR Tickets (Ridership),QR %,NCMC (Ridership),NCMC %,Total Passenger Flow\n"
        '1,Feb-24,"10,000",50%,"9,000",45%,"1,000",5%,"20,000"\n'
        '2,Mar-24,"11,000",50%,"10,000",45%,"1,000",5%,"22,000"\n', encoding="utf-8")
    import docx
    d = docx.Document()
    d.add_heading("Station accessibility policy", 1)
    d.add_paragraph("Wheelchair ramps must be provided at every station entrance.")
    d.save(docs / "Accessibility_Policy.docx")
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assets"
    ws.append(["asset", "station", "count"])
    ws.append(["Lift", "Egmore", 4])
    ws.append(["Escalator", "Egmore", 6])
    wb.save(docs / "asset_counts.xlsx")


_build_raw()

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.main import app, bootstrap  # noqa: E402
from app.services.jobs import enqueue, run_pending_sync  # noqa: E402
from app.db import SessionLocal  # noqa: E402

app.state.testing = True


@pytest.fixture(scope="session")
def client():
    with TestClient(app) as c:
        db = SessionLocal()
        enqueue(db, "INGEST_DATASETS", {})
        enqueue(db, "INGEST_DOCUMENTS", {})
        db.close()
        run_pending_sync()
        yield c


def _login(client, email, password):
    r = client.post("/api/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="session")
def admin(client):
    return _login(client, "admin@megaasset.local", "ChangeMe!2026")


@pytest.fixture(scope="session")
def operator(client):
    return _login(client, "operator@megaasset.local", "Operator!2026")


@pytest.fixture(scope="session")
def supervisor(client):
    return _login(client, "supervisor@megaasset.local", "Supervisor!2026")


@pytest.fixture(scope="session")
def engineer(client):
    return _login(client, "engineer@megaasset.local", "Engineer!2026")


def make_passenger(client, email):
    r = client.post("/api/auth/signup", json={"full_name": "Test Passenger", "email": email, "password": "Passw0rd!x"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="session")
def passenger(client):
    return make_passenger(client, "passenger1@example.com")


@pytest.fixture(scope="session")
def station_id(client, admin):
    items = client.get("/api/stations", headers=admin).json()["items"]
    return next(s["id"] for s in items if s["code"] == "SEG")  # Egmore


PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00"
       b"\x0cIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82")


def raise_complaint(client, headers, station_id, category="lifts_escalators", text="The escalator near gate A2 is not working since morning", files=None, **extra):
    r = client.post("/api/complaints", headers=headers, data={"description": text, "station_id": station_id, "category": category, "language": "en", **extra}, files=files or [])
    assert r.status_code == 200, r.text
    return r.json()
