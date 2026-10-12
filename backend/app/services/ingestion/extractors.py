"""Text/table extraction. Runs in worker processes, so everything here is picklable and stateless.

Each extractor returns a dict:
  {"units": [{"page": int|None, "label": str|None, "text": str, "ocr": bool,
              "rows": [(row_number, row_text), ...] | None}],
   "page_count": int, "ocr_pages": int, "empty_pages": int, "warnings": [str]}
"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

SUPPORTED = {".pdf", ".txt", ".md", ".markdown", ".csv", ".xlsx", ".docx", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}

_ocr_engine = None


def _ocr():
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR  # optional dependency
        _ocr_engine = RapidOCR()
    return _ocr_engine


def _ocr_image_bytes(data: bytes) -> str:
    result, _ = _ocr()(data, use_cls=False)
    if not result:
        return ""
    # result: [[box, text, score], ...] ; order top-to-bottom, left-to-right
    items = sorted(result, key=lambda r: (round(r[0][0][1] / 12), r[0][0][0]))
    lines, last_y, cur = [], None, []
    for box, text, score in items:
        y = round(box[0][1] / 12)
        if last_y is not None and y != last_y and cur:
            lines.append(" ".join(cur)); cur = []
        cur.append(text); last_y = y
    if cur:
        lines.append(" ".join(cur))
    return "\n".join(lines)


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("­", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def ocr_pdf_pages(path: str, pages: list[int], dpi: int = 150) -> dict[int, str]:
    """OCR selected 1-based pages of a PDF. Used by the background OCR job."""
    import pymupdf

    out = {}
    with pymupdf.open(path) as doc:
        for p in pages:
            if 1 <= p <= len(doc):
                pix = doc[p - 1].get_pixmap(dpi=dpi)
                out[p] = _clean(_ocr_image_bytes(pix.tobytes("png")))
    return out


def extract_pdf(path: str, enable_ocr: bool = False) -> dict:
    """Text-layer extraction. Pages without a text layer are flagged in `needs_ocr` and handled by the OCR job
    (or OCR'd inline when enable_ocr=True)."""
    import pymupdf

    units, warnings, needs_ocr = [], [], []
    ocr_pages = empty = 0
    with pymupdf.open(path) as doc:
        n = len(doc)
        for i in range(n):
            page = doc[i]
            blocks = page.get_text("blocks", sort=True)
            paras = [_clean(b[4]) for b in blocks if b[6] == 0 and _clean(b[4])]
            text = "\n\n".join(paras)
            used_ocr = False
            if len(text) < 30 and enable_ocr:
                try:
                    pix = page.get_pixmap(dpi=170)
                    text = _clean(_ocr_image_bytes(pix.tobytes("png")))
                    used_ocr = bool(text)
                    if used_ocr:
                        ocr_pages += 1
                except Exception as exc:  # OCR is optional
                    warnings.append(f"OCR unavailable for page {i + 1}: {exc}")
            if len(text) < 30:
                empty += 1
                if not enable_ocr:
                    needs_ocr.append(i + 1)
            units.append({"page": i + 1, "label": None, "text": text, "ocr": used_ocr, "rows": None})
    return {"units": units, "page_count": n, "ocr_pages": ocr_pages, "empty_pages": empty, "warnings": warnings,
            "needs_ocr": needs_ocr}


def extract_text(path: str) -> dict:
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    text = _clean(text.replace("\r\n", "\n"))
    return {"units": [{"page": 1, "label": None, "text": text, "ocr": False, "rows": None}],
            "page_count": 1, "ocr_pages": 0, "empty_pages": 0 if text else 1, "warnings": []}


def _row_text(header, row) -> str:
    parts = []
    for h, v in zip(header, row):
        if v is None or str(v).strip() == "":
            continue
        parts.append(f"{h}: {v}")
    return "; ".join(parts)


def extract_csv(path: str) -> dict:
    raw = Path(path).read_bytes().decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(raw))
    rows = list(reader)
    if not rows:
        return {"units": [], "page_count": 0, "ocr_pages": 0, "empty_pages": 1, "warnings": ["Empty CSV"]}
    header = [h.strip() or f"col{i + 1}" for i, h in enumerate(rows[0])]
    # Spreadsheet row numbers: header is row 1, first data row is row 2
    data = [(idx + 2, _row_text(header, r)) for idx, r in enumerate(rows[1:]) if any(c.strip() for c in r)]
    text = "Columns: " + ", ".join(header) + "\n" + "\n".join(t for _, t in data)
    return {"units": [{"page": 1, "label": Path(path).stem, "text": text, "ocr": False, "rows": data,
                       "header": header}],
            "page_count": 1, "ocr_pages": 0, "empty_pages": 0, "warnings": []}


def extract_xlsx(path: str) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    units = []
    for si, ws in enumerate(wb.worksheets, start=1):
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip() if h is not None else f"col{i + 1}" for i, h in enumerate(rows[0])]
        data = []
        for ridx, r in enumerate(rows[1:], start=2):
            vals = ["" if v is None else (v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else v) for v in r]
            if any(str(v).strip() for v in vals):
                data.append((ridx, _row_text(header, vals)))
        text = f"Sheet: {ws.title}\nColumns: " + ", ".join(header) + "\n" + "\n".join(t for _, t in data)
        units.append({"page": si, "label": ws.title, "text": text, "ocr": False, "rows": data, "header": header})
    wb.close()
    return {"units": units, "page_count": len(units), "ocr_pages": 0, "empty_pages": 0, "warnings": []}


def extract_docx(path: str) -> dict:
    import docx

    d = docx.Document(path)
    lines = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style == "title":
            lines.append(f"\n# {t}\n")
        else:
            lines.append(t)
    for ti, table in enumerate(d.tables, start=1):
        lines.append(f"\n# Table {ti}\n")
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    text = _clean("\n".join(lines))
    # DOCX has no fixed pagination; we cite by section heading instead of page.
    return {"units": [{"page": None, "label": None, "text": text, "ocr": False, "rows": None}],
            "page_count": 0, "ocr_pages": 0, "empty_pages": 0 if text else 1, "warnings": []}


def extract_image(path: str, enable_ocr: bool = True) -> dict:
    if not enable_ocr:
        return {"units": [], "page_count": 1, "ocr_pages": 0, "empty_pages": 1,
                "warnings": ["OCR disabled; image text not extracted"]}
    text = _clean(_ocr_image_bytes(Path(path).read_bytes()))
    return {"units": [{"page": 1, "label": None, "text": text, "ocr": True, "rows": None}],
            "page_count": 1, "ocr_pages": 1 if text else 0, "empty_pages": 0 if text else 1, "warnings": []}


def extract(path: str, enable_ocr: bool = True) -> dict:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path, enable_ocr=False)  # scanned pages are deferred to the OCR job
    if ext in (".txt", ".md", ".markdown"):
        return extract_text(path)
    if ext == ".csv":
        return extract_csv(path)
    if ext == ".xlsx":
        return extract_xlsx(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
        return extract_image(path, enable_ocr)
    raise ValueError(f"Unsupported file type: {ext}")


def safe_extract(path: str, enable_ocr: bool = True) -> dict:
    """Process-pool entry point: never raises, returns {'error': ...} instead."""
    try:
        return extract(path, enable_ocr)
    except Exception as exc:  # pragma: no cover - surfaced to the job record
        return {"error": f"{type(exc).__name__}: {exc}"}
