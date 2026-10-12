"""Structured dataset processing: CMRL GTFS, Chennai-wide GTFS (kept separate), ridership CSV, passenger-flow XLSX.

Stations come only from CMRL GTFS parent-station records (location_type=1). Access assets (entrances, and lifts,
escalators and ramps referenced in entrance descriptions) are DOCUMENTED with the exact stops.txt row as source.
"""
from __future__ import annotations

import csv
import logging
import re
from collections import Counter, defaultdict
from datetime import datetime

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from ..config import settings
from ..models import Asset, Dataset, HourlyRidership, RidershipMonthly, Station, StationRidership, utcnow
from .ingestion.pipeline import rel_path, sha256_file

log = logging.getLogger("megaasset.datasets")

# Curated aliases used for matching station names in other sources. Each alias maps to a GTFS zone code.
STATION_ALIASES = {
    "SCC": ["CENTRAL METRO", "CHENNAI CENTRAL", "MGR CENTRAL", "M.G.R. CENTRAL", "M.G.R CENTRAL",
            "PURATCHI THALAIVAR DR. M.G. RAMACHANDRAN CENTRAL", "PURATCHI THALAIVAR DR.M.G.RAMACHANDRAN CENTRAL"],
    "SCM": ["CMBT", "PURATCHI THALAIVI DR. J. JAYALALITHAA CMBT", "PURATCHI THALAIVI DR.J.JAYALALITHAA CMBT"],
    "SAL": ["ALANDUR", "ARIGNAR ANNA ALANDUR"],
    "SPC": ["PACHAIYAPPA'S COLLEGE", "PACHAIYAPPAS COLLEGE", "PACHAIYAPPA"],
    "SMM": ["ST. THOMAS MOUNT", "ST THOMAS MOUNT", "ST.THOMAS MOUNT"],
    "SOT": ["NANGANALLUR ROAD", "OTA - NANGANALLUR ROAD", "OTA NANGANALLUR ROAD"],
    "STL": ["THOUSAND LIGHT", "THOUSAND LIGHTS"],
    "SWD": ["WIMCO NAGAR DEPOT"],
    "SSI": ["EKKATTUTHANGAL", "EKKATUTHANGAL"],
    "SGM": ["AG-DMS", "AG DMS"],
}
# Station names that are also common words/places; require an explicit "metro"/"station" context in documents.
AMBIGUOUS = {"AIRPORT", "LIC", "HIGH COURT", "CENTRAL", "GUINDY", "KOYAMBEDU", "EGMORE", "SAIDAPET", "VADAPALANI",
             "GOVERNMENT ESTATE", "MEENAMBAKKAM", "WASHERMENPET", "TONDIARPET", "THIRUVOTRIYUR"}


def norm_name(s: str) -> str:
    s = s.upper().replace("&", " AND ")
    s = re.sub(r"\b(METRO|STATION|STN|RLY|RAILWAY)\b", " ", s)
    s = re.sub(r"[^A-Z0-9' ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def display_name(gtfs_name: str) -> str:
    n = gtfs_name.strip()
    special = {"AG-DMS": "AG-DMS", "LIC": "LIC"}
    if n in special:
        return special[n]
    t = n.title().replace("Dr.M.G.Ramachandran", "Dr. M.G. Ramachandran").replace("Dr.J.Jayalalithaa", "Dr. J. Jayalalithaa")
    t = t.replace("St.Thomas", "St. Thomas").replace("Cmbt", "CMBT").replace("Ag-Dms", "AG-DMS")
    return t


def _read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def _upsert_dataset(db: Session, slug, **kw) -> Dataset:
    ds = db.scalar(select(Dataset).where(Dataset.slug == slug))
    if not ds:
        ds = Dataset(slug=slug, **kw)
        db.add(ds)
    else:
        for k, v in kw.items():
            setattr(ds, k, v)
    return ds


def _gtfs_summary(folder) -> tuple[list, int, dict]:
    tables, total = [], 0
    for p in sorted(folder.glob("*.txt")):
        with open(p, encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            header = next(reader, [])
            n = sum(1 for _ in reader)
        tables.append({"name": p.name, "columns": header, "rows": n})
        total += n
    return tables, total, {}


def load_cmrl_gtfs(db: Session) -> Dataset:
    folder = settings.raw_data_dir / "cmrl" / "gtfs"
    stops = _read_csv(folder / "stops.txt")
    routes = {r["route_id"]: r for r in _read_csv(folder / "routes.txt")}
    trips = {t["trip_id"]: t for t in _read_csv(folder / "trips.txt")}
    stop_times = _read_csv(folder / "stop_times.txt")

    stations = [s for s in stops if s.get("location_type", "").strip() == "1"]
    entrances = [s for s in stops if s.get("location_type", "").strip() == "0"]
    parent_of = {e["stop_id"]: e["parent_station"].strip() for e in entrances}

    # Lines per station, from GTFS route colours (#000092 = Blue, #00A700 = Green) via stop_times -> trips -> routes
    colour_line = {"000092": "Blue Line", "00A700": "Green Line"}
    station_lines = defaultdict(set)
    station_routes = defaultdict(set)
    trips_per_route = Counter()
    for t in trips.values():
        trips_per_route[t["route_id"]] += 1
    for st in stop_times:
        trip = trips.get(st["trip_id"])
        if not trip:
            continue
        parent = parent_of.get(st["stop_id"])
        route = routes.get(trip["route_id"], {})
        line = colour_line.get(route.get("route_color", "").strip().upper())
        if parent:
            station_routes[parent].add(route.get("route_short_name", "").strip())
            if line:
                station_lines[parent].add(line)

    by_code = {}
    for s in stations:
        code = s["zone_id"].strip()
        st = db.scalar(select(Station).where(Station.code == code))
        name = s["stop_name"].strip()
        aliases = sorted({norm_name(name), *[norm_name(a) for a in STATION_ALIASES.get(code, [])]})
        fields = dict(gtfs_stop_id=s["stop_id"].strip(), name=name, display_name=display_name(name),
                      lat=float(s["stop_lat"]), lon=float(s["stop_lon"]), lines=sorted(station_lines.get(s["stop_id"].strip(), [])),
                      aliases=aliases, source=f"data/raw/cmrl/gtfs/stops.txt (stop_id {s['stop_id'].strip()})",
                      provenance="DOCUMENTED")
        if st:
            for k, v in fields.items():
                setattr(st, k, v)
        else:
            st = Station(code=code, **fields)
            db.add(st)
        by_code[s["stop_id"].strip()] = st
    db.flush()

    # Access assets from entrance records
    existing = {a.asset_code: a for a in db.scalars(select(Asset).where(Asset.source.like("%gtfs/stops.txt%")))}
    keep = set()
    for row_no, e in enumerate(stops, start=2):
        if e.get("location_type", "").strip() != "0":
            continue
        st = by_code.get(e["parent_station"].strip())
        desc = (e.get("stop_desc") or "").strip()
        if not st or not desc:
            continue
        label = desc.split(" ENTRANCE")[0].strip()
        u = desc.upper()
        modes = [m for m, kw in (("lift", "LIFT"), ("escalator", "ESCALATOR"), ("ramp", "RAMP"), ("staircase", "STAIRCASE")) if kw in u]
        src = f"data/raw/cmrl/gtfs/stops.txt row {row_no} (stop_id {e['stop_id'].strip()})"
        items = [("ENT", "station_entrance", f"Entrance {label} — {st.display_name}")]
        if "LIFT" in u:
            items.append(("LIFT", "lift", f"Lift at/near entrance {label} — {st.display_name}"))
        if "ESCALATOR" in u:
            items.append(("ESC", "escalator", f"Escalator near entrance {label} — {st.display_name}"))
        if "RAMP" in u:
            items.append(("RAMP", "accessibility_ramp", f"Accessibility ramp at entrance {label} — {st.display_name}"))
        for suffix, atype, name in items:
            code = f"{st.code}-{suffix}-{e['stop_id'].strip()}"
            keep.add(code)
            fields = dict(asset_type=atype, name=name, station_id=st.id, description=desc.title(),
                          lat=float(e["stop_lat"]), lon=float(e["stop_lon"]), source=src, source_ref=e["stop_id"].strip(),
                          provenance="DOCUMENTED",
                          attributes={"entrance": label, "access_modes": modes, "gtfs_stop_desc": desc,
                                      "platform_code": (e.get("platform_code") or "").strip(),
                                      "note": "Referenced in GTFS entrance description; condition and maintenance state are not recorded in this source."})
            a = existing.get(code)
            if a:
                for k, v in fields.items():
                    setattr(a, k, v)
            else:
                db.add(Asset(asset_code=code, **fields))
    for code, a in existing.items():
        if code not in keep:
            db.delete(a)

    # Service statistics (PROCESSED from stop_times)
    route_stats = []
    for rid, r in routes.items():
        route_stats.append({"route_id": rid, "short_name": r.get("route_short_name", "").strip(),
                            "long_name": r.get("route_long_name", "").strip(),
                            "line": colour_line.get(r.get("route_color", "").strip().upper(), "Unclassified (no colour in feed)"),
                            "trips": trips_per_route.get(rid, 0)})
    first_dep = defaultdict(lambda: "99:99:99")
    last_dep = defaultdict(lambda: "00:00:00")
    deps_per_station = Counter()
    for st in stop_times:
        p = parent_of.get(st["stop_id"])
        if not p:
            continue
        deps_per_station[p] += 1
        d = st["departure_time"]
        first_dep[p] = min(first_dep[p], d)
        last_dep[p] = max(last_dep[p], d)
    station_service = {by_code[p].code: {"departures": deps_per_station[p], "first": first_dep[p], "last": last_dep[p],
                                         "routes": sorted(station_routes[p])}
                       for p in by_code if p in deps_per_station}
    cal = _read_csv(folder / "calendar.txt")
    tables, total, _ = _gtfs_summary(folder)
    ds = _upsert_dataset(
        db, "cmrl-gtfs", name="CMRL GTFS feed", scope="CMRL_ONLY", kind="GTFS",
        description="Chennai Metro Rail Limited static GTFS: stations, entrances, routes, trips and timetables.",
        rel_path=rel_path(folder), checksum=sha256_file(folder / "stops.txt"), status="PROCESSED", error=None,
        row_count=total, tables=tables, provenance="DOCUMENTED", processed_at=utcnow(),
        stats={"stations": len(stations), "entrances": len(entrances), "routes": len(routes), "trips": len(trips),
               "stop_times": len(stop_times), "route_stats": route_stats, "station_service": station_service,
               "calendar": cal, "line_colours": {"#000092": "Blue Line", "#00A700": "Green Line"},
               "note": "Feed validity per calendar.txt; timetable may not reflect current operations."})
    return ds


def load_chennai_gtfs(db: Session) -> Dataset | None:
    folder = settings.raw_data_dir / "chennai_mobility"
    if not (folder / "stops.txt").exists():
        return None
    agencies = _read_csv(folder / "agency.txt")
    tables, total, _ = _gtfs_summary(folder)
    routes_by_agency = Counter(r.get("agency_id", "") for r in _read_csv(folder / "routes.txt"))
    return _upsert_dataset(
        db, "chennai-gtfs", name="Chennai-wide multimodal GTFS", scope="CHENNAI_WIDE", kind="GTFS",
        description="City-wide GTFS (MTC bus, CMRL and other agencies). Kept separate from CMRL-only records; not used to derive CMRL stations or assets.",
        rel_path=rel_path(folder), checksum=sha256_file(folder / "agency.txt"), status="PROCESSED", error=None,
        row_count=total, tables=tables, provenance="DOCUMENTED", processed_at=utcnow(),
        stats={"agencies": [{"id": a["agency_id"], "name": a["agency_name"], "routes": routes_by_agency.get(a["agency_id"], 0)} for a in agencies]})


def _to_int(v):
    if v is None:
        return None
    s = str(v).replace(",", "").strip()
    try:
        return int(float(s))
    except ValueError:
        return None


def load_monthly_ridership(db: Session) -> Dataset | None:
    path = settings.raw_data_dir / "cmrl" / "ridership" / "cmrl_monthly_ridership.csv"
    if not path.exists():
        return None
    rows = _read_csv(path)
    db.execute(delete(RidershipMonthly))
    parsed = 0
    for i, r in enumerate(rows, start=2):
        label = (r.get("Month") or "").strip()
        try:
            month_start = datetime.strptime(label, "%b-%y").strftime("%Y-%m-01")
        except ValueError:
            continue
        db.add(RidershipMonthly(month_label=label, month_start=month_start, closed_loop=_to_int(r.get("Closed Loop (Ridership)")),
                                qr=_to_int(r.get("QR Tickets (Ridership)")), ncmc=_to_int(r.get("NCMC (Ridership)")),
                                total=_to_int(r.get("Total Passenger Flow")), source_row=i))
        parsed += 1
    header = list(rows[0].keys()) if rows else []
    return _upsert_dataset(
        db, "cmrl-monthly-ridership", name="CMRL monthly ridership", scope="CMRL_ONLY", kind="CSV",
        description="Monthly passenger flow by ticket media (closed loop card, QR, NCMC).",
        rel_path=rel_path(path), checksum=sha256_file(path), status="PROCESSED", error=None, row_count=len(rows),
        tables=[{"name": path.name, "columns": header, "rows": len(rows)}], provenance="DOCUMENTED",
        processed_at=utcnow(), stats={"parsed_rows": parsed, "skipped_rows": len(rows) - parsed})


def _match_station(name: str, stations: list[Station]) -> tuple[Station | None, str | None]:
    n = norm_name(name)
    for s in stations:
        if norm_name(s.name) == n:
            return s, "exact"
    for s in stations:
        if n in (s.aliases or []):
            return s, "alias"
    for s in stations:
        sn = norm_name(s.name)
        if sn and (sn in n or n in sn) and min(len(sn), len(n)) >= 4:
            return s, "contains"
    return None, None


def load_passenger_flow(db: Session) -> Dataset | None:
    import openpyxl

    path = settings.raw_data_dir / "cmrl" / "passenger_flow" / "chennai_metro_phase1_ridership.xlsx"
    if not path.exists():
        return None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    stations = list(db.scalars(select(Station)))
    db.execute(delete(StationRidership))
    db.execute(delete(HourlyRidership))
    tables, unmatched, matched = [], set(), 0
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        tables.append({"name": ws.title, "columns": [str(c) for c in rows[0]] if rows else [], "rows": max(0, len(rows) - 1)})
        if ws.title == "Station simple":
            h = {k: i for i, k in enumerate(rows[0])}
            for rn, r in enumerate(rows[1:], start=2):
                st, method = _match_station(str(r[h["station_name"]]), stations)
                if st:
                    matched += 1
                else:
                    unmatched.add(str(r[h["station_name"]]))
                db.add(StationRidership(service_date=r[h["service_date"]].strftime("%Y-%m-%d"), day_label=str(r[h["day_label"]]),
                                        is_partial="partial" in str(r[h["day_label"]]), line_name=str(r[h["line_name"]]),
                                        station_seq=int(r[h["station_seq"]]), station_name=str(r[h["station_name"]]),
                                        station_id=st.id if st else None, match_method=method,
                                        passengers=int(r[h["passengers"]] or 0), source_row=rn))
        elif ws.title == "Hourly network":
            h = {k: i for i, k in enumerate(rows[0])}
            for rn, r in enumerate(rows[1:], start=2):
                db.add(HourlyRidership(service_date=r[h["service_date"]].strftime("%Y-%m-%d"),
                                       is_partial=bool(r[h["is_partial_day"]]), hour=int(r[h["hour"]]),
                                       total=int(r[h["total_passengers"]] or 0), source_row=rn))
    summary = []
    ws = wb["Summary"]
    rows = list(ws.iter_rows(values_only=True))
    for r in rows[1:]:
        summary.append({str(k): (v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else v) for k, v in zip(rows[0], r)})
    wb.close()
    return _upsert_dataset(
        db, "cmrl-passenger-flow", name="CMRL Phase 1 station passenger flow", scope="CMRL_ONLY", kind="XLSX",
        description="Station-level and hourly network passenger counts for the service dates contained in the workbook.",
        rel_path=rel_path(path), checksum=sha256_file(path), status="PROCESSED", error=None,
        row_count=sum(t["rows"] for t in tables), tables=tables, provenance="DOCUMENTED", processed_at=utcnow(),
        stats={"summary": summary, "station_rows_matched": matched, "unmatched_station_names": sorted(unmatched),
               "note": "Covers only the service dates in the file; a partial day is flagged and excluded from totals."})


def load_all_datasets(db: Session) -> dict:
    out = {}
    for fn in (load_cmrl_gtfs, load_chennai_gtfs, load_monthly_ridership, load_passenger_flow):
        try:
            ds = fn(db)
            db.commit()
            if ds:
                out[ds.slug] = ds.status
        except Exception as exc:
            db.rollback()
            log.exception("Dataset load failed: %s", fn.__name__)
            out[fn.__name__] = f"FAILED: {exc}"
    return out
