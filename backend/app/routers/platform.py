"""Overview, analytics, datasets, AI model status, notifications, global search and administration."""
from __future__ import annotations

import csv
from collections import Counter, defaultdict
from datetime import timedelta
from itertools import islice

from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from ..agents.workflow import CATEGORY_KEYWORDS, STAGES
from ..config import PROJECT_DIR, settings
from ..db import get_db
from ..models import (Asset, AuditLog, Complaint, Dataset, Document, DocumentChunk, GraphEntity, GraphRelationship,
                      Investigation, Job, Notification, ResolutionMemory, Role, SolutionProposal, Station, Ticket,
                      TicketEvent, User, utcnow)
from ..provenance import ALL_LABELS, COMPLAINT_CATEGORIES, TEAMS
from ..security import ROLES, get_current_user, has_perm, require_admin, require_staff
from ..serializers import document_dict, iso, station_brief, user_dict
from ..services.common import audit, user_brief
from ..services.embeddings import embedder_status, read_meta, vector_index
from ..services.jobs import enqueue
from ..services.llm import llm_status
from ..services.retrieval import reranker_status
from ..services.tickets import policy

router = APIRouter(prefix="/api", tags=["platform"])


@router.get("/overview")
def overview(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    staff = has_perm(user, "ticket:read_all")
    now = utcnow()
    tq = select(Ticket).join(Complaint)
    if not staff:
        tq = tq.where(Complaint.user_id == user.id)
    tickets = list(db.scalars(tq))
    status_counts = Counter(t.status for t in tickets)
    open_ = [t for t in tickets if t.status not in ("RESOLVED", "CLOSED")]
    week = [t for t in tickets if t.created_at >= now - timedelta(days=7)]
    awaiting = [t for t in tickets if t.status == "AWAITING_VERIFICATION"]
    by_cat = Counter(t.complaint.category for t in tickets)
    by_station = Counter(t.complaint.station.display_name for t in tickets)
    pri = Counter(t.priority for t in open_)
    docs_indexed = db.scalar(select(func.count(Document.id)).where(Document.status.in_(["INDEXED", "PARTIAL"]))) or 0
    docs_total = db.scalar(select(func.count(Document.id))) or 0
    chunks = db.scalar(select(func.count(DocumentChunk.id))) or 0
    pages = db.scalar(select(func.sum(Document.page_count)).where(Document.status.in_(["INDEXED", "PARTIAL"]))) or 0
    daily = defaultdict(int)
    for t in tickets:
        if t.created_at >= now - timedelta(days=30):
            daily[t.created_at.strftime("%Y-%m-%d")] += 1
    series = [{"date": (now - timedelta(days=i)).strftime("%Y-%m-%d"), "count": daily.get((now - timedelta(days=i)).strftime("%Y-%m-%d"), 0)} for i in range(29, -1, -1)]
    ev_q = select(TicketEvent).order_by(TicketEvent.id.desc()).limit(12)
    if not staff:
        ev_q = select(TicketEvent).join(Ticket).join(Complaint).where(Complaint.user_id == user.id).order_by(TicketEvent.id.desc()).limit(12)
    activity = [{"id": e.id, "event_type": e.event_type, "ticket_id": e.ticket_id, "ticket_code": e.ticket.ticket_code, "note": e.note,
                 "to_status": e.to_status, "actor": user_brief(e.actor), "created_at": iso(e.created_at)} for e in db.scalars(ev_q)]
    from ..models import RidershipMonthly
    rid = [{"month": r.month_label, "total": r.total} for r in db.scalars(select(RidershipMonthly).order_by(RidershipMonthly.month_start))]
    if not staff:
        resolved = sum(1 for t in tickets if t.status in ("RESOLVED", "CLOSED"))
        return {"scope": "mine",
                "tickets": {"total": len(tickets), "open": len(open_), "this_week": len(week), "awaiting_verification": len(awaiting),
                            "resolved": resolved, "by_status": status_counts,
                            "by_category": {COMPLAINT_CATEGORIES[k]["en"]: v for k, v in by_cat.items()},
                            "by_station": dict(by_station.most_common(8)), "open_by_priority": pri, "daily": series},
                "activity": activity}
    out = {
        "scope": "all" if staff else "mine",
        "tickets": {"total": len(tickets), "open": len(open_), "this_week": len(week), "awaiting_verification": len(awaiting),
                    "by_status": status_counts, "by_category": {COMPLAINT_CATEGORIES[k]["en"]: v for k, v in by_cat.items()},
                    "by_station": dict(by_station.most_common(8)), "open_by_priority": pri, "daily": series},
        "documents": {"indexed": docs_indexed, "total": docs_total, "chunks": chunks, "pages": pages},
        "stations": db.scalar(select(func.count(Station.id))) or 0,
        "assets": db.scalar(select(func.count(Asset.id))) or 0,
        "graph": {"entities": db.scalar(select(func.count(GraphEntity.id))) or 0, "relationships": db.scalar(select(func.count(GraphRelationship.id))) or 0},
        "investigations": db.scalar(select(func.count(Investigation.id))) or 0 if staff else None,
        "pending_approvals": db.scalar(select(func.count(SolutionProposal.id)).where(SolutionProposal.status == "PROPOSED")) or 0 if staff else None,
        "resolutions": dict(db.execute(select(ResolutionMemory.outcome, func.count()).group_by(ResolutionMemory.outcome)).all()) if staff else None,
        "ridership": rid,
        "activity": activity,
    }
    return out


@router.get("/analytics")
def analytics(user: User = Depends(require_staff), db: Session = Depends(get_db)):
    tickets = list(db.scalars(select(Ticket)))
    res_hours = [((t.resolved_at - t.created_at).total_seconds() / 3600) for t in tickets if t.resolved_at]
    cat_station = defaultdict(lambda: defaultdict(int))
    for t in tickets:
        cat_station[t.complaint.station.display_name][COMPLAINT_CATEGORIES[t.complaint.category]["en"]] += 1
    events = list(db.scalars(select(TicketEvent).where(TicketEvent.to_status.isnot(None)).order_by(TicketEvent.ticket_id, TicketEvent.id)))
    time_in = defaultdict(list)
    last = {}
    for e in events:
        if e.ticket_id in last:
            prev_status, prev_time = last[e.ticket_id]
            time_in[prev_status].append((e.created_at - prev_time).total_seconds() / 3600)
        last[e.ticket_id] = (e.to_status, e.created_at)
    lang = Counter(t.complaint.language for t in tickets)
    teams = Counter(t.assigned_team or "Unassigned" for t in tickets if t.status not in ("RESOLVED", "CLOSED"))
    verif = Counter()
    from ..models import VerificationRecord
    for v in db.scalars(select(VerificationRecord)):
        verif[v.outcome] += 1
    docs = [{"category": c, "documents": n, "chunks": ch or 0, "pages": p or 0, "ocr_pages": o or 0} for c, n, ch, p, o in db.execute(
        select(Document.category, func.count(), func.sum(Document.chunk_count), func.sum(Document.page_count), func.sum(Document.ocr_pages))
        .where(Document.status.in_(["INDEXED", "PARTIAL"])).group_by(Document.category))]
    return {
        "ticket_count": len(tickets),
        "resolution_hours": {"count": len(res_hours), "median": sorted(res_hours)[len(res_hours) // 2] if res_hours else None,
                             "mean": sum(res_hours) / len(res_hours) if res_hours else None},
        "time_in_status_hours": {k: round(sum(v) / len(v), 2) for k, v in time_in.items() if v},
        "station_category": [{"station": s, **cats, "total": sum(cats.values())} for s, cats in sorted(cat_station.items(), key=lambda x: -sum(x[1].values()))[:12]],
        "languages": lang, "open_by_team": teams, "verification_outcomes": verif,
        "reopen_rate": round(sum(1 for t in tickets if t.reopen_count) / len(tickets), 3) if tickets else None,
        "documents": docs,
        "provenance_note": "All figures are computed from records in the MegaAsset database (PROCESSED). Empty values mean no records exist yet.",
    }


# ---------------- Datasets ----------------
@router.get("/datasets")
def datasets(user: User = Depends(require_staff), db: Session = Depends(get_db)):
    out = []
    for d in db.scalars(select(Dataset).order_by(Dataset.scope, Dataset.name)):
        stats = {k: v for k, v in (d.stats or {}).items() if k not in ("route_stats", "station_service", "summary")}
        out.append({"id": d.id, "slug": d.slug, "name": d.name, "description": d.description, "scope": d.scope, "kind": d.kind,
                    "rel_path": d.rel_path, "checksum": d.checksum, "status": d.status, "error": d.error, "row_count": d.row_count,
                    "tables": d.tables, "stats": stats, "provenance": d.provenance, "processed_at": iso(d.processed_at)})
    return out


@router.get("/datasets/{slug}/preview")
def dataset_preview(slug: str, table: str | None = None, limit: int = 25, user: User = Depends(require_staff), db: Session = Depends(get_db)):
    d = db.scalar(select(Dataset).where(Dataset.slug == slug))
    if not d:
        raise HTTPException(404, "Dataset not found")
    base = (PROJECT_DIR / d.rel_path).resolve()
    if settings.raw_data_dir.resolve() not in base.parents and base != settings.raw_data_dir.resolve():
        raise HTTPException(400, "Invalid dataset path")
    limit = min(limit, 100)
    if d.kind == "GTFS":
        names = [t["name"] for t in d.tables]
        table = table if table in names else names[0]
        path = base / table
        with open(path, encoding="utf-8-sig", newline="") as f:
            r = csv.reader(f)
            header = next(r, [])
            rows = list(islice(r, limit))
        return {"table": table, "tables": names, "columns": header, "rows": rows}
    if d.kind == "CSV":
        with open(base, encoding="utf-8-sig", newline="") as f:
            r = csv.reader(f)
            header = next(r, [])
            rows = list(islice(r, limit))
        return {"table": base.name, "tables": [base.name], "columns": header, "rows": rows}
    import openpyxl
    wb = openpyxl.load_workbook(base, read_only=True, data_only=True)
    names = wb.sheetnames
    table = table if table in names else names[0]
    ws = wb[table]
    it = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(it, [])]
    rows = [[(v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else v) for v in r] for r in islice(it, limit)]
    wb.close()
    return {"table": table, "tables": names, "columns": header, "rows": rows}


@router.post("/datasets/reprocess")
def reprocess(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    job = enqueue(db, "INGEST_DATASETS", {}, user.id)
    return {"job_id": job.id}


# ---------------- AI models ----------------
@router.get("/models")
def models(user: User = Depends(require_staff), db: Session = Depends(get_db)):
    try:
        import rapidocr_onnxruntime  # noqa: F401
        ocr = {"available": True, "engine": "RapidOCR (ONNX Runtime, local)", "enabled": settings.enable_ocr}
    except Exception as exc:
        ocr = {"available": False, "engine": "RapidOCR", "enabled": settings.enable_ocr, "error": str(exc),
               "setup": "pip install rapidocr_onnxruntime"}
    emb = embedder_status()
    meta = read_meta()
    embedded = db.scalar(select(func.count(DocumentChunk.id)).where(DocumentChunk.embedding.isnot(None))) or 0
    return {
        "embedding": {**emb, "index_model": meta.get("embedding_model"), "index_dim": meta.get("embedding_dim"),
                      "embedded_chunks": embedded, "vector_index_size": vector_index.size, "provenance": "PROCESSED",
                      "description": "Dense passage embeddings for semantic retrieval; stored per chunk in SQLite and loaded as a normalised matrix."},
        "keyword_index": {"engine": "SQLite FTS5 (BM25, Porter stemming)", "available": True},
        "reranker": {**reranker_status(), "available": reranker_status()["loaded"],
                     "note": "Candidates from reciprocal-rank fusion (BM25 + vector) are re-ordered by a local cross-encoder for direct questions; "
                             "broad multi-document questions use fusion ranking with per-document diversity."},
        "workflow_policy": policy(),
        "llm": llm_status(),
        "ocr": ocr,
        "classifier": {"type": "Transparent keyword rules (English + Tamil)", "categories": {k: v[:8] for k, v in CATEGORY_KEYWORDS.items()},
                       "provenance": "AI_HYPOTHESIS", "evaluation": "Not a trained model; no accuracy metric is claimed."},
        "risk_scoring": {"type": "Transparent weighted rules over records", "provenance": "POTENTIAL",
                         "evaluation": "No labelled outcome dataset exists in the workspace; predictive accuracy is not claimed."},
        "workflow": {"engine": "LangGraph StateGraph", "stages": STAGES, "human_approval": True},
        "provenance_labels": ALL_LABELS,
    }


# ---------------- Notifications ----------------
@router.get("/notifications")
def notifications(unread_only: bool = False, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    q = select(Notification).where(Notification.user_id == user.id)
    if unread_only:
        q = q.where(Notification.is_read.is_(False))
    rows = db.scalars(q.order_by(Notification.id.desc()).limit(50))
    unread = db.scalar(select(func.count(Notification.id)).where(Notification.user_id == user.id, Notification.is_read.is_(False))) or 0
    return {"unread": unread, "items": [{"id": n.id, "type": n.type, "title": n.title, "body": n.body, "link": n.link, "is_read": n.is_read,
                                         "created_at": iso(n.created_at)} for n in rows]}


@router.post("/notifications/{nid}/read")
def mark_read(nid: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    n = db.get(Notification, nid)
    if not n or n.user_id != user.id:
        raise HTTPException(404, "Notification not found")
    n.is_read = True
    db.commit()
    return {"ok": True}


@router.post("/notifications/read-all")
def mark_all(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    for n in db.scalars(select(Notification).where(Notification.user_id == user.id, Notification.is_read.is_(False))):
        n.is_read = True
    db.commit()
    return {"ok": True}


# ---------------- Global search ----------------
@router.get("/search")
def search(q: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    q = q.strip()
    if len(q) < 2:
        return {"results": []}
    like = f"%{q}%"
    res = []
    staff = has_perm(user, "ticket:read_all")
    for s in (db.scalars(select(Station).where(or_(Station.display_name.ilike(like), Station.code.ilike(like))).limit(5)) if staff else []):
        res.append({"type": "station", "id": s.id, "title": s.display_name, "subtitle": f"{s.code} · {', '.join(s.lines) or 'line n/a'}", "link": f"/app/stations/{s.id}"})
    tq = select(Ticket).join(Complaint).where(or_(Ticket.ticket_code.ilike(like), Complaint.complaint_code.ilike(like), Complaint.description.ilike(like)))
    if not has_perm(user, "ticket:read_all"):
        tq = tq.where(Complaint.user_id == user.id)
    for t in db.scalars(tq.limit(6)):
        link = f"/app/tickets/{t.id}" if has_perm(user, "ticket:read_all") else f"/app/complaints/{t.complaint_id}"
        res.append({"type": "ticket", "id": t.id, "title": f"{t.ticket_code} · {t.complaint.complaint_code}", "subtitle": f"{t.status} · {t.complaint.description[:70]}", "link": link})
    for d in (db.scalars(select(Document).where(or_(Document.title.ilike(like), Document.filename.ilike(like))).limit(5)) if staff else []):
        res.append({"type": "document", "id": d.id, "title": d.title, "subtitle": f"{d.category} · {d.status}", "link": f"/app/documents?doc={d.id}"})
    for a in (db.scalars(select(Asset).where(or_(Asset.name.ilike(like), Asset.asset_code.ilike(like))).limit(5)) if staff else []):
        res.append({"type": "asset", "id": a.id, "title": a.name, "subtitle": a.asset_code, "link": f"/app/assets/{a.id}"})
    if has_perm(user, "ticket:read_all"):
        for i in db.scalars(select(Investigation).where(or_(Investigation.code.ilike(like), Investigation.question.ilike(like))).limit(4)):
            res.append({"type": "investigation", "id": i.id, "title": i.code, "subtitle": i.question[:80], "link": f"/app/investigations/{i.id}"})
    return {"results": res}


# ---------------- Administration ----------------
@router.get("/admin/users")
def users(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    return {"items": [user_dict(u) for u in db.scalars(select(User).order_by(User.id))],
            "roles": [{"name": k, **v} for k, v in ROLES.items()], "teams": TEAMS}


class UserPatch(BaseModel):
    role: str | None = None
    team: str | None = None
    is_active: bool | None = None


@router.patch("/admin/users/{uid}")
def patch_user(uid: int, body: UserPatch, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    u = db.get(User, uid)
    if not u:
        raise HTTPException(404, "User not found")
    if u.id == user.id and (body.role not in (None, "admin") or body.is_active is False):
        raise HTTPException(409, "You cannot remove your own administrator access")
    if body.role:
        if body.role not in ROLES:
            raise HTTPException(422, "Unknown role")
        u.role = body.role
    if body.team is not None:
        if body.team and body.team not in TEAMS:
            raise HTTPException(422, "Unknown team")
        u.team = body.team or None
    if body.is_active is not None:
        u.is_active = body.is_active
    audit(db, user, "user.update", "user", u.id, **body.model_dump(exclude_none=True))
    db.commit()
    return user_dict(u)


@router.get("/admin/audit")
def audit_log(limit: int = 100, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    return [{"id": a.id, "action": a.action, "entity_type": a.entity_type, "entity_id": a.entity_id, "detail": a.detail,
             "actor": user_brief(a.actor), "created_at": iso(a.created_at)} for a in db.scalars(select(AuditLog).order_by(AuditLog.id.desc()).limit(min(limit, 500)))]


def _job_dict(j: Job, full=False):
    d = {"id": j.id, "type": j.type, "status": j.status, "progress": j.progress, "total": j.total, "message": j.message,
         "params": j.params, "error": j.error, "created_at": iso(j.created_at), "started_at": iso(j.started_at),
         "finished_at": iso(j.finished_at), "cancel_requested": j.cancel_requested}
    if full:
        d["result"] = {k: v for k, v in (j.result or {}).items() if k != "traceback"}
    return d


@router.get("/jobs")
def jobs(user: User = Depends(require_staff), db: Session = Depends(get_db)):
    return [_job_dict(j) for j in db.scalars(select(Job).order_by(Job.id.desc()).limit(50))]


@router.get("/jobs/{jid}")
def job(jid: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    j = db.get(Job, jid)
    if not j or (j.created_by != user.id and not has_perm(user, "ticket:read_all")):
        raise HTTPException(404, "Job not found")
    return _job_dict(j, full=True)


@router.post("/jobs/{jid}/cancel")
def cancel_job(jid: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    j = db.get(Job, jid)
    if not j or (j.created_by != user.id and user.role != "admin"):
        raise HTTPException(404, "Job not found")
    if j.status == "QUEUED":
        j.status, j.message, j.finished_at = "CANCELLED", "Cancelled before start", utcnow()
    elif j.status == "RUNNING":
        j.cancel_requested = True
        j.message = "Cancellation requested"
    db.commit()
    return _job_dict(j)


@router.get("/system")
def system(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    import os
    import platform as pf
    db_size = settings.database_path.stat().st_size if settings.database_path.exists() else 0
    return {"database_path": str(settings.database_path), "database_bytes": db_size, "raw_data_dir": str(settings.raw_data_dir),
            "uploads_dir": str(settings.uploads_dir), "python": pf.python_version(), "cpu_count": os.cpu_count(),
            "ingest_workers": settings.ingest_workers, "embedding_batch_size": settings.embedding_batch_size,
            "ocr_enabled": settings.enable_ocr, "email_sms": "Not configured — notifications are in-app only.",
            "tables": {t: db.execute(__import__("sqlalchemy").text(f"SELECT COUNT(*) FROM {t}")).scalar() for t in
                       ["users", "stations", "assets", "documents", "document_chunks", "complaints", "tickets", "ticket_events",
                        "investigations", "solution_proposals", "resolution_memory", "notifications", "graph_entities",
                        "graph_relationships", "audit_logs", "jobs"]}}
